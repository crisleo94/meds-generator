import io
import re
import openpyxl
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.graphics.barcode import createBarcodeDrawing

from reportlab.lib import colors

# --- Dimensiones exactas ---
CARD_W = 3.5 * 72 / 2.54   # ~99.2126 pt (3.5 cm)
CARD_H = 1.5 * 72 / 2.54   # ~42.5197 pt (1.5 cm)

PAGE_W, PAGE_H = A4        # 595.28 pt x 841.89 pt

COLS = 5
ROWS = 18
CARDS_PER_PAGE = COLS * ROWS  # 90 tarjetas por página

MARGIN_X = (PAGE_W - COLS * CARD_W) / 2.0
MARGIN_Y = (PAGE_H - ROWS * CARD_H) / 2.0

MAX_TEXT_W = CARD_W - 8.0  # ~91.21 pt
MAX_TEXT_H = CARD_H - 5.0  # ~37.52 pt
MAX_BC_W = CARD_W - 10.0   # ~89.21 pt

def wrap_text_to_width(text, font_name, font_size, max_w):
    """
    Divide un texto en líneas de forma estricta según el ancho máximo permitido,
    partiendo palabras largas por caracteres si fuera estrictamente necesario.
    """
    words = text.split()
    if not words:
        return []
    lines = []
    current_line = []

    for word in words:
        # Si la palabra individual supera el ancho máximo, la dividimos por caracteres
        if stringWidth(word, font_name, font_size) > max_w:
            if current_line:
                lines.append(' '.join(current_line))
                current_line = []
            chunk = ""
            for char in word:
                if stringWidth(chunk + char, font_name, font_size) <= max_w:
                    chunk += char
                else:
                    if chunk:
                        lines.append(chunk)
                    chunk = char
            if chunk:
                current_line = [chunk]
            continue

        test_line = ' '.join(current_line + [word]) if current_line else word
        if stringWidth(test_line, font_name, font_size) <= max_w:
            current_line.append(word)
        else:
            if current_line:
                lines.append(' '.join(current_line))
            current_line = [word]

    if current_line:
        lines.append(' '.join(current_line))
    return lines

def balance_wrap_2(words, font_name, font_size, max_w):
    """Divide las palabras en 2 líneas equilibradas evitando palabras o números huérfanos."""
    n = len(words)
    best_split = None
    min_cost = float('inf')
    for i in range(1, n):
        l1 = ' '.join(words[:i])
        l2 = ' '.join(words[i:])
        w1 = stringWidth(l1, font_name, font_size)
        w2 = stringWidth(l2, font_name, font_size)
        if w1 <= max_w and w2 <= max_w:
            diff = abs(w1 - w2)
            cost = diff
            if len(words[i:]) == 1:
                cost += 40.0
                if len(words[i:][0]) <= 3:
                    cost += 100.0
            if len(words[:i]) == 1 and len(words[0]) <= 3:
                cost += 80.0
            if cost < min_cost:
                min_cost = cost
                best_split = [l1, l2]
    return best_split

def balance_wrap_3(words, font_name, font_size, max_w):
    """Divide las palabras en 3 líneas en caso de textos extremos."""
    n = len(words)
    best_split = None
    min_cost = float('inf')
    for i in range(1, n - 1):
        for j in range(i + 1, n):
            l1 = ' '.join(words[:i])
            l2 = ' '.join(words[i:j])
            l3 = ' '.join(words[j:])
            w1 = stringWidth(l1, font_name, font_size)
            w2 = stringWidth(l2, font_name, font_size)
            w3 = stringWidth(l3, font_name, font_size)
            if w1 <= max_w and w2 <= max_w and w3 <= max_w:
                cost = max(w1, w2, w3) - min(w1, w2, w3)
                if len(words[j:]) == 1 and len(words[j:][0]) <= 3:
                    cost += 100.0
                if cost < min_cost:
                    min_cost = cost
                    best_split = [l1, l2, l3]
    return best_split

def layout_single_column_text(desc, font_name='Helvetica', custom_font_size=None):
    """Layout para tarjeta de 1 columna: Texto grande centrado horizontal y verticalmente."""
    words = desc.split()

    # Si el usuario eligió un tamaño de fuente fijo (ej. 14 o 16 pt), se respeta estrictamente
    if custom_font_size is not None and custom_font_size > 0:
        fs = float(custom_font_size)
        lines = wrap_text_to_width(desc, font_name, fs, MAX_TEXT_W)
        if not lines:
            lines = [desc]
        leading = fs * 1.15
        total_h = (len(lines) - 1) * leading + fs
        return {
            'lines': lines,
            'font': font_name,
            'font_size': fs,
            'leading': leading,
            'total_h': total_h
        }

    # Modo Auto (Adaptativo)
    # 1. Probar en 1 sola línea (fuentes de 12.0 bajando a 7.0)
    for fs in [12.0, 11.0, 10.0, 9.5, 9.0, 8.5, 8.0, 7.5, 7.0]:
        w = stringWidth(desc, font_name, fs)
        if w <= MAX_TEXT_W:
            return {
                'lines': [desc],
                'font': font_name,
                'font_size': fs,
                'leading': fs * 1.2,
                'total_h': fs
            }

    # 2. Probar en 2 líneas equilibradas (fuentes de 10.0 bajando a 6.0)
    for fs in [10.0, 9.5, 9.0, 8.5, 8.0, 7.5, 7.0, 6.5, 6.0]:
        lines = balance_wrap_2(words, font_name, fs, MAX_TEXT_W)
        if lines:
            leading = fs * 1.15
            total_h = leading + fs
            if total_h <= MAX_TEXT_H:
                return {
                    'lines': lines,
                    'font': font_name,
                    'font_size': fs,
                    'leading': leading,
                    'total_h': total_h
                }

    # 3. Probar en 3 líneas (fuentes de 8.0 bajando a 5.0)
    for fs in [8.0, 7.5, 7.0, 6.5, 6.0, 5.5, 5.0]:
        lines = balance_wrap_3(words, font_name, fs, MAX_TEXT_W)
        if lines:
            leading = fs * 1.12
            total_h = 2 * leading + fs
            if total_h <= MAX_TEXT_H:
                return {
                    'lines': lines,
                    'font': font_name,
                    'font_size': fs,
                    'leading': leading,
                    'total_h': total_h
                }

    lines = wrap_text_to_width(desc, font_name, 5.0, MAX_TEXT_W)
    return {
        'lines': lines if lines else [desc],
        'font': font_name,
        'font_size': 5.0,
        'leading': 6.0,
        'total_h': (len(lines) - 1) * 6.0 + 5.0 if lines else 16.0
    }

def layout_single_column_barcode(code):
    """Layout para tarjeta de 1 columna: Código de barras grande centrado con número en negrita."""
    bc_h = 16.0
    cve_fs = 6.0
    gap = 2.0
    total_h = bc_h + gap + cve_fs
    return {
        'bc_h': bc_h,
        'cve_fs': cve_fs,
        'gap': gap,
        'total_h': total_h
    }

def layout_card_numbers_only(desc, cve, font_name='Helvetica', custom_font_size=None):
    """Layout para tarjetas de 2 columnas sin código de barras (solo número en negrita)."""
    title_font = font_name
    cve_font = 'Helvetica-Bold'
    words = desc.split()

    if custom_font_size is not None and custom_font_size > 0:
        fs = float(custom_font_size)
        lines = wrap_text_to_width(desc, title_font, fs, MAX_TEXT_W)
        if not lines:
            lines = [desc]
        leading = fs * 1.15
        cve_fs = min(fs - 1.5, 6.5)
        cve_w = stringWidth(cve, cve_font, cve_fs)
        if cve_w > MAX_TEXT_W:
            cve_fs *= (MAX_TEXT_W / cve_w)
        gap = 2.5
        h_total = (len(lines) - 1) * leading + fs + gap + cve_fs
        return {
            'title_font': title_font,
            'title_fs': fs,
            'title_leading': leading,
            'title_lines': lines,
            'cve_font': cve_font,
            'cve_fs': cve_fs,
            'gap': gap,
            'h_total': h_total
        }

    for fs in [8.0, 7.5, 7.0]:
        w = stringWidth(desc, title_font, fs)
        if w <= MAX_TEXT_W:
            cve_fs = min(fs - 1.8, 6.0)
            cve_w = stringWidth(cve, cve_font, cve_fs)
            if cve_w > MAX_TEXT_W:
                cve_fs *= (MAX_TEXT_W / cve_w)
            gap = 3.5
            h_total = fs + gap + cve_fs
            return {
                'title_font': title_font,
                'title_fs': fs,
                'title_leading': fs * 1.15,
                'title_lines': [desc],
                'cve_font': cve_font,
                'cve_fs': cve_fs,
                'gap': gap,
                'h_total': h_total
            }

    for fs in [8.0, 7.5, 7.0, 6.5, 6.0]:
        lines = balance_wrap_2(words, title_font, fs, MAX_TEXT_W)
        if lines:
            leading = fs * 1.15
            cve_fs = min(fs - 1.8, 5.8)
            cve_w = stringWidth(cve, cve_font, cve_fs)
            if cve_w > MAX_TEXT_W:
                cve_fs *= (MAX_TEXT_W / cve_w)
            gap = 2.8
            h_total = leading + fs + gap + cve_fs
            if h_total <= MAX_TEXT_H:
                return {
                    'title_font': title_font,
                    'title_fs': fs,
                    'title_leading': leading,
                    'title_lines': lines,
                    'cve_font': cve_font,
                    'cve_fs': cve_fs,
                    'gap': gap,
                    'h_total': h_total
                }

    for fs in [6.0, 5.5, 5.0]:
        lines = balance_wrap_3(words, title_font, fs, MAX_TEXT_W)
        if lines:
            leading = fs * 1.12
            cve_fs = min(fs - 1.2, 5.0)
            gap = 2.0
            h_total = 2 * leading + fs + gap + cve_fs
            return {
                'title_font': title_font,
                'title_fs': fs,
                'title_leading': leading,
                'title_lines': lines,
                'cve_font': cve_font,
                'cve_fs': cve_fs,
                'gap': gap,
                'h_total': h_total
            }

    lines = wrap_text_to_width(desc, title_font, 5.0, MAX_TEXT_W)
    return {
        'title_font': title_font,
        'title_fs': 5.0,
        'title_leading': 5.5,
        'title_lines': lines if lines else [desc],
        'cve_font': cve_font,
        'cve_fs': 4.5,
        'gap': 2.0,
        'h_total': 15.0
    }

def layout_barcode_card(desc, cve, font_name='Helvetica', custom_font_size=None):
    """Layout para tarjetas de 2 columnas con código de barras y número debajo."""
    title_font = font_name
    words = desc.split()

    if custom_font_size is not None and custom_font_size > 0:
        fs = float(custom_font_size)
        lines = wrap_text_to_width(desc, title_font, fs, MAX_TEXT_W)
        if not lines:
            lines = [desc]
        leading = fs * 1.15
        t_h = (len(lines) - 1) * leading + fs
        bc_h = 9.0
        cve_fs = 5.0
        gap1 = 2.0
        gap2 = 1.5
        total_h = t_h + gap1 + bc_h + gap2 + cve_fs
        return {
            't_lines': lines,
            't_fs': fs,
            't_leading': leading,
            'bc_h': bc_h,
            'cve_fs': cve_fs,
            'gap1': gap1,
            'gap2': gap2,
            'total_h': total_h
        }

    for fs in [7.5, 7.0]:
        w = stringWidth(desc, title_font, fs)
        if w <= MAX_TEXT_W:
            t_lines = [desc]
            t_fs = fs
            t_leading = fs * 1.15
            t_h = fs
            bc_h = 11.0
            cve_fs = 5.4
            gap1 = 2.8
            gap2 = 1.8
            total_h = t_h + gap1 + bc_h + gap2 + cve_fs
            return {
                't_lines': t_lines,
                't_fs': t_fs,
                't_leading': t_leading,
                'bc_h': bc_h,
                'cve_fs': cve_fs,
                'gap1': gap1,
                'gap2': gap2,
                'total_h': total_h
            }

    for fs in [7.0, 6.5, 6.0, 5.5]:
        lines = balance_wrap_2(words, title_font, fs, MAX_TEXT_W)
        if lines:
            t_lines = lines
            t_fs = fs
            t_leading = fs * 1.12
            t_h = t_leading + fs
            bc_h = 9.2
            cve_fs = 5.0
            gap1 = 2.0
            gap2 = 1.5
            total_h = t_h + gap1 + bc_h + gap2 + cve_fs
            return {
                't_lines': t_lines,
                't_fs': t_fs,
                't_leading': t_leading,
                'bc_h': bc_h,
                'cve_fs': cve_fs,
                'gap1': gap1,
                'gap2': gap2,
                'total_h': total_h
            }

    lines = wrap_text_to_width(desc, title_font, 5.0, MAX_TEXT_W)
    return {
        't_lines': lines if lines else [desc],
        't_fs': 5.0,
        't_leading': 5.5,
        'bc_h': 8.0,
        'cve_fs': 4.5,
        'gap1': 1.5,
        'gap2': 1.5,
        'total_h': 24.0
    }

def clean_value(raw_val, remove_symbols=True):
    if raw_val is None:
        return ""
    val_str = str(raw_val).strip()
    if remove_symbols:
        val_str = re.sub(r'^[*/\s]+', '', val_str).strip()
    return val_str

def extract_excel_info(file_path):
    """Lee el archivo Excel y devuelve las hojas y las columnas de la primera hoja."""
    wb = openpyxl.load_workbook(file_path, read_only=True)
    sheet_names = wb.sheetnames
    first_sheet = wb[sheet_names[0]]
    rows_iter = first_sheet.iter_rows(values_only=True)
    header_row = next(rows_iter, None)

    headers = []
    if header_row:
        for idx, h in enumerate(header_row):
            h_clean = str(h).strip() if h is not None else f"Columna_{idx+1}"
            headers.append(h_clean)

    row_count = 0
    sample_rows = []
    for r in rows_iter:
        if any(cell is not None and str(cell).strip() != '' for cell in r):
            row_count += 1
            if len(sample_rows) < 5:
                sample_rows.append([str(c).strip() if c is not None else '' for c in r])

    wb.close()
    return {
        'sheets': sheet_names,
        'headers': headers,
        'row_count': row_count,
        'sample_rows': sample_rows
    }

def generate_pdf_buffer(file_path, col1, col2=None, num_cols=2, card_style='barcode',
                        clean_symbols=True, max_items=None,
                        font_family='Helvetica', font_size='auto', font_color='#000000'):
    """
    Genera el PDF en memoria soportando 1 columna o 2 columnas con tipografía, tamaño y color configurables:
    - num_cols = 1:
        - card_style = 'text_only': Texto grande centrado
        - card_style = 'barcode_only': Código de barras grande centrado
    - num_cols = 2:
        - card_style = 'barcode': Título + Código de barras + número
        - card_style = 'numbers': Título + solo número en negrita
    """
    # Parsear tamaño de fuente personalizado si no es 'auto'
    custom_font_size = None
    if font_size and str(font_size).strip().lower() != 'auto':
        try:
            custom_font_size = float(font_size)
        except ValueError:
            custom_font_size = None

    # Parsear color de fuente
    try:
        text_color = colors.HexColor(font_color)
    except Exception:
        text_color = colors.HexColor('#000000')

    wb = openpyxl.load_workbook(file_path, read_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)

    if not header_row:
        wb.close()
        raise ValueError("El archivo Excel está vacío o no contiene encabezados.")

    headers = [str(h).strip() if h is not None else '' for h in header_row]

    col1_idx = headers.index(col1) if col1 in headers else 0
    col2_idx = headers.index(col2) if col2 and col2 in headers else (1 if len(headers) > 1 else 0)

    items = []
    for r in rows_iter:
        if r is None:
            continue
        val1 = r[col1_idx] if col1_idx < len(r) else ''
        val2 = r[col2_idx] if (num_cols == 2 and col2_idx < len(r)) else ''

        clean_val1 = clean_value(val1, remove_symbols=clean_symbols)
        clean_val2 = str(val2).strip() if val2 is not None else ''

        if num_cols == 1:
            if clean_val1:
                items.append((clean_val1, ''))
        else:
            if clean_val1 or clean_val2:
                items.append((clean_val1, clean_val2))

        if max_items and len(items) >= max_items:
            break

    wb.close()

    total_items = len(items)
    total_pages = (total_items + CARDS_PER_PAGE - 1) // CARDS_PER_PAGE
    if total_pages == 0:
        total_pages = 1

    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=A4)

    item_idx = 0
    for page_num in range(1, total_pages + 1):
        for r_idx in range(ROWS):
            for c_idx in range(COLS):
                if item_idx >= total_items:
                    break
                v1, v2 = items[item_idx]
                item_idx += 1

                x0 = MARGIN_X + c_idx * CARD_W
                y0 = MARGIN_Y + (ROWS - 1 - r_idx) * CARD_H

                # Borde sutil gris claro para corte
                c.setStrokeColorRGB(0.78, 0.78, 0.78)
                c.setLineWidth(0.4)
                c.rect(x0, y0, CARD_W, CARD_H)

                # Máscara de recorte estricta dentro de la tarjeta
                c.saveState()
                clip_p = c.beginPath()
                clip_p.rect(x0 + 0.5, y0 + 0.5, CARD_W - 1.0, CARD_H - 1.0)
                c.clipPath(clip_p, stroke=0, fill=0)

                # ==========================================
                # CASO 1: UNA SOLA COLUMNA
                # ==========================================
                if num_cols == 1:
                    if card_style == 'barcode_only' and v1:
                        # Código de barras grande centrado
                        layout = layout_single_column_barcode(v1)
                        bc_h = layout['bc_h']
                        cve_fs = layout['cve_fs']
                        gap = layout['gap']
                        total_h = layout['total_h']

                        y_top = y0 + (CARD_H + total_h) / 2.0
                        bc_y = y_top - bc_h

                        bw = 0.44
                        bc = createBarcodeDrawing('Code128', value=v1, barWidth=bw, barHeight=bc_h, humanReadable=False)
                        if bc.width > MAX_BC_W:
                            bw = bw * (MAX_BC_W / bc.width)
                            bc = createBarcodeDrawing('Code128', value=v1, barWidth=bw, barHeight=bc_h, humanReadable=False)

                        bc_x = x0 + (CARD_W - bc.width) / 2.0
                        bc.drawOn(c, bc_x, bc_y)

                        y_cve_baseline = bc_y - gap - 0.75 * cve_fs
                        c.setFillColor(text_color)
                        c.setFont('Helvetica-Bold', cve_fs)
                        cve_w = stringWidth(v1, 'Helvetica-Bold', cve_fs)
                        c.drawString(x0 + (CARD_W - cve_w) / 2.0, y_cve_baseline, v1)
                    else:
                        # Texto grande centrado
                        layout = layout_single_column_text(v1, font_name=font_family, custom_font_size=custom_font_size)
                        lines = layout['lines']
                        fs = layout['font_size']
                        leading = layout['leading']
                        total_h = layout['total_h']

                        y_top = y0 + (CARD_H + total_h) / 2.0
                        y_first_baseline = y_top - 0.75 * fs

                        c.setFillColor(text_color)
                        c.setFont(layout['font'], fs)
                        for l_idx, line in enumerate(lines):
                            line_y = y_first_baseline - l_idx * leading
                            line_w = stringWidth(line, layout['font'], fs)
                            c.drawString(x0 + (CARD_W - line_w) / 2.0, line_y, line)

                # ==========================================
                # CASO 2: DOS COLUMNAS
                # ==========================================
                else:
                    if card_style == 'barcode' and v2:
                        # Modo Código de Barras (v1 = título, v2 = código)
                        layout = layout_barcode_card(v1, v2, font_name=font_family, custom_font_size=custom_font_size)
                        t_lines = layout['t_lines']
                        t_fs = layout['t_fs']
                        t_leading = layout['t_leading']
                        bc_h = layout['bc_h']
                        cve_fs = layout['cve_fs']
                        gap1 = layout['gap1']
                        gap2 = layout['gap2']
                        total_h = layout['total_h']

                        y_top = y0 + (CARD_H + total_h) / 2.0

                        # Título
                        c.setFillColor(text_color)
                        c.setFont(font_family, t_fs)
                        y_first_baseline = y_top - 0.75 * t_fs
                        for l_idx, line in enumerate(t_lines):
                            line_y = y_first_baseline - l_idx * t_leading
                            line_w = stringWidth(line, font_family, t_fs)
                            c.drawString(x0 + (CARD_W - line_w) / 2.0, line_y, line)

                        y_last_baseline = y_first_baseline - (len(t_lines) - 1) * t_leading
                        bc_y = y_last_baseline - 0.25 * t_fs - gap1 - bc_h

                        bw = 0.40
                        bc = createBarcodeDrawing('Code128', value=v2, barWidth=bw, barHeight=bc_h, humanReadable=False)
                        if bc.width > MAX_BC_W:
                            bw = bw * (MAX_BC_W / bc.width)
                            bc = createBarcodeDrawing('Code128', value=v2, barWidth=bw, barHeight=bc_h, humanReadable=False)

                        bc_x = x0 + (CARD_W - bc.width) / 2.0
                        bc.drawOn(c, bc_x, bc_y)

                        y_cve_baseline = bc_y - gap2 - 0.75 * cve_fs
                        c.setFillColor(text_color)
                        c.setFont('Helvetica-Bold', cve_fs)
                        cve_w = stringWidth(v2, 'Helvetica-Bold', cve_fs)
                        c.drawString(x0 + (CARD_W - cve_w) / 2.0, y_cve_baseline, v2)

                    else:
                        # Modo Solo Números
                        layout = layout_card_numbers_only(v1, v2, font_name=font_family, custom_font_size=custom_font_size)
                        h_total = layout['h_total']
                        t_fs = layout['title_fs']
                        t_leading = layout['title_leading']
                        t_lines = layout['title_lines']
                        cve_fs = layout['cve_fs']
                        gap = layout['gap']

                        y_top = y0 + (CARD_H + h_total) / 2.0
                        y_first_baseline = y_top - 0.75 * t_fs

                        c.setFillColor(text_color)
                        c.setFont(layout['title_font'], t_fs)
                        for l_idx, line in enumerate(t_lines):
                            line_y = y_first_baseline - l_idx * t_leading
                            line_w = stringWidth(line, layout['title_font'], t_fs)
                            c.drawString(x0 + (CARD_W - line_w) / 2.0, line_y, line)

                        y_last_baseline = y_first_baseline - (len(t_lines) - 1) * t_leading
                        y_cve_baseline = y_last_baseline - 0.25 * t_fs - 0.75 * cve_fs - gap

                        c.setFont(layout['cve_font'], cve_fs)
                        cve_w = stringWidth(v2, layout['cve_font'], cve_fs)
                        c.drawString(x0 + (CARD_W - cve_w) / 2.0, y_cve_baseline, v2)

                c.restoreState()

            if item_idx >= total_items:
                break

        c.showPage()

    c.save()
    buffer.seek(0)
    return buffer, total_items, total_pages
