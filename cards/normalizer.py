import io
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from .sorter import clean_and_sort_lines

def normalize_excel_columns(file_path, selected_cols, code_col=None, remove_duplicates=True):
    """
    Consolida múltiples columnas de medicamentos en una sola lista estructurada,
    limpia prefijos/viñetas, ordena alfabéticamente y genera un nuevo archivo Excel.
    """
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows:
        raise ValueError("El archivo Excel no contiene datos.")

    header_row = rows[0]
    headers = [str(h).strip() if h is not None else f"Col_{i+1}" for i, h in enumerate(header_row)]

    # Mapear índices de columnas a consolidar
    consolidate_indices = []
    for col_name in selected_cols:
        if col_name in headers:
            consolidate_indices.append(headers.index(col_name))

    if not consolidate_indices:
        raise ValueError("No se seleccionó ninguna columna válida para consolidar.")

    code_idx = headers.index(code_col) if code_col and code_col in headers else None

    # Recolectar pares (medicamento, codigo)
    raw_items = []
    for r in rows[1:]:
        if not r:
            continue
        cve_val = str(r[code_idx]).strip() if code_idx is not None and code_idx < len(r) and r[code_idx] is not None else ""
        for idx in consolidate_indices:
            if idx < len(r) and r[idx] is not None:
                val = str(r[idx]).strip()
                # Limpieza de bordes: viñetas, comillas, guiones, asteriscos
                clean_val = re.sub(r'^[-*•\s"\'`]+|["\'`\s]+$', '', val)
                clean_val = re.sub(r'\s+', ' ', clean_val).strip()
                if clean_val:
                    raw_items.append((clean_val, cve_val))

    # Ordenar y/o deduplicar
    if remove_duplicates:
        seen = set()
        deduped = []
        for desc, cve in raw_items:
            key = desc.lower()
            if key not in seen:
                seen.add(key)
                deduped.append((desc, cve))
        raw_items = deduped

    # Ordenamiento alfabético estricto
    # Reutilizamos la función de ordenamiento
    dummy_text = '\n'.join(item[0] for item in raw_items)
    sort_info = clean_and_sort_lines(dummy_text, remove_duplicates=False)
    sorted_names = sort_info['items']

    # Asignar de nuevo el orden a los pares
    name_to_cve = {desc: cve for desc, cve in raw_items}
    final_items = [(name, name_to_cve.get(name, "")) for name in sorted_names]

    # Crear nuevo libro Excel con diseño profesional
    out_wb = openpyxl.Workbook()
    out_ws = out_wb.active
    out_ws.title = "NORMALIZADO"

    # Estilos de cabecera
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    out_headers = ["DESC"]
    if code_col:
        out_headers.append("CVEART")

    out_ws.append(out_headers)

    for cell in out_ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border

    out_ws.row_dimensions[1].height = 24

    for r_idx, (desc, cve) in enumerate(final_items, start=2):
        row_data = [desc]
        if code_col:
            row_data.append(cve)
        out_ws.append(row_data)

        out_ws.cell(row=r_idx, column=1).alignment = Alignment(horizontal="left", vertical="center")
        out_ws.cell(row=r_idx, column=1).border = thin_border
        if code_col:
            out_ws.cell(row=r_idx, column=2).alignment = Alignment(horizontal="center", vertical="center")
            out_ws.cell(row=r_idx, column=2).border = thin_border
        out_ws.row_dimensions[r_idx].height = 18

    # Autoajuste de anchos de columna
    for col in out_ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        out_ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    buffer = io.BytesIO()
    out_wb.save(buffer)
    out_wb.close()
    buffer.seek(0)

    return {
        'buffer': buffer,
        'total_items': len(final_items),
        'sample_items': final_items[:6]
    }
